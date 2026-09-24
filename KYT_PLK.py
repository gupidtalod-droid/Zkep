import datetime
from openpyxl import Workbook
import streamlit as st

st.title("📋 ระบบบันทึก Morning KYT & Warehouse Readiness")
st.write(
    "กรอกข้อมูลประจำวัน สามารถกดปุ่มเพิ่มหัวข้อย่อย ประเด็น หรือปัญหาได้ตามต้องการ"
)

# กำหนดค่าเริ่มต้นใน session_state
if "main_count" not in st.session_state:
  st.session_state.main_count = 2
if "ready_count" not in st.session_state:
  st.session_state.ready_count = 4
if "kyt_count" not in st.session_state:
  st.session_state.kyt_count = 3
if "problem_count" not in st.session_state:
  st.session_state.problem_count = 1  # เริ่มต้นที่ 1 ข้อ (ไม่มี)
if "file_ready" not in st.session_state:
  st.session_state.file_ready = False

with st.form("kyt_form"):
  st.subheader("📌 ข้อมูลทั่วไป")
  col1, col2 = st.columns(2)
  with col1:
    branch = st.text_input("สาขา", "พิษณุโลก")
    reporter = st.text_input("ผู้รายงาน", "วิไลวรรณ ศิริแสน")
  with col2:
    # ปรับใช้ปฏิทินเลือกวันที่จริง
    selected_date = st.date_input("วันที่", datetime.date.today())
    date_str = selected_date.strftime("%d/%m/%Y")

    time_str = st.text_input("ประชุมทีมเสร็จเวลา", "08:10 น.")

  participants = st.text_input("ผู้เข้าร่วม", "เช็กเกอร์และพนักงานรายวัน")

  # --- 1. งานหลักวันนี้ ---
  st.subheader("1. งานหลักวันนี้")
  main_works = []
  default_mains = [
      "รับสินค้า 5 เที่ยว / จ่ายสินค้า 15 เที่ยว",
      "เตรียมสินค้าจัดส่งก่อน 10:00 น. 8 เที่ยว",
  ]
  for i in range(st.session_state.main_count):
    default_val = (
        default_mains[i] if i < len(default_mains) else f"งานหลักข้อที่ {i+1}"
    )
    val = st.text_input(f"งานหลักข้อที่ {i+1}", default_val)
    main_works.append(val)

  col_btn1, _ = st.columns([1, 4])
  with col_btn1:
    if st.form_submit_button("➕ เพิ่มงานหลัก"):
      st.session_state.main_count += 1
      st.rerun()

  # --- 2. ความพร้อมก่อนเริ่มงาน ---
  st.subheader("2. ความพร้อมก่อนเริ่มงาน")
  ready_items = []
  default_readies = [
      "พนักงานมา 12 คน ครบตามแผน",
      (
          "พนักงานรายวันมา 13 จากแผน 14 คน จัดคนงานแยกตามจุดลงสินค้าแต่ละประเภทแล้ว"
      ),
      "รถโฟล์คลิฟท์ 1 คัน พร้อมใช้งาน",
      "พื้นที่รับ–จ่ายสินค้า พร้อมใช้งาน",
  ]
  for i in range(st.session_state.ready_count):
    default_val = (
        default_readies[i]
        if i < len(default_readies)
        else f"ความพร้อมข้อที่ {i+1}"
    )
    val = st.text_input(f"ความพร้อมข้อที่ {i+1}", default_val)
    ready_items.append(val)

  col_btn2, _ = st.columns([1, 4])
  with col_btn2:
    if st.form_submit_button("➕ เพิ่มความพร้อม"):
      st.session_state.ready_count += 1
      st.rerun()

  # --- 3. ประเด็น KYT ที่คุยกับทีมวันนี้ ---
  st.subheader("3. ประเด็น KYT ที่คุยกับทีมวันนี้")
  kyt_items = []
  default_kyts = [
      "เน้นย้ำการเช็กงานเตรียมรอกระจาย",
      "เน้นย้ำดูแลการเข้า–ออกบริเวณจุดโหลด",
      (
          "ย้ำการทำ 5ส ให้พื้นที่ทำงานและจุดเก็บเอกสารสะอาด เป็นระเบียบ"
          " ค้นหาของได้รวดเร็ว"
      ),
  ]
  for i in range(st.session_state.kyt_count):
    default_val = (
        default_kyts[i] if i < len(default_kyts) else f"ประเด็น KYT ข้อที่ {i+1}"
    )
    val = st.text_input(f"ประเด็น KYT ข้อที่ {i+1}", default_val)
    kyt_items.append(val)

  col_btn3, _ = st.columns([1, 4])
  with col_btn3:
    if st.form_submit_button("➕ เพิ่มประเด็น KYT"):
      st.session_state.kyt_count += 1
      st.rerun()

  # --- 4. ปัญหาค้าง/เรื่องที่ต้องประสาน ---
  st.subheader("4. ปัญหาค้าง/เรื่องที่ต้องประสาน")
  problem_items = []
  default_problems = ["ไม่มี"]
  for i in range(st.session_state.problem_count):
    default_val = (
        default_problems[i]
        if i < len(default_problems)
        else f"ปัญหาข้อที่ {i+1}"
    )
    val = st.text_input(f"ปัญหาข้อที่ {i+1}", default_val)
    problem_items.append(val)

  col_btn4, _ = st.columns([1, 4])
  with col_btn4:
    if st.form_submit_button("➕ เพิ่มปัญหา"):
      st.session_state.problem_count += 1
      st.rerun()

  # --- 5. สถานะสาขา ---
  st.subheader("5. สถานะสาขา")
  status = st.selectbox(
      "สถานะ",
      ["ดำเนินงานได้แต่มีข้อจำกัด", "ดำเนินงานได้ปกติ", "มีปัญหาติดขัด"],
  )
  status_detail = st.text_area(
      "รายละเอียดสถานะ",
      (
          "งานเช้าดำเนินได้ตามแผนที่ปรับแล้ว"
          " ยังไม่กระทบเวลารับ–จ่ายสินค้า"
      ),
  )

  # ปุ่มกดบันทึกหลัก
  submitted = st.form_submit_button("💾 บันทึกข้อมูลเป็น Excel ทั้งหมด")

  if submitted:
    wb = Workbook()
    ws = wb.active
    ws.title = "Morning KYT"

    data = [
        ["Morning KYT & Warehouse Readiness", ""],
        ["สาขา:", branch],
        ["วันที่:", date_str],
        ["ผู้รายงาน:", reporter],
        ["ประชุมทีมเสร็จเวลา:", time_str],
        ["ผู้เข้าร่วม:", participants],
        ["---", "---"],
        ["1. งานหลักวันนี้", ""],
    ]

    for item in main_works:
      data.append(["-", item])

    data.append(["---", "---"])
    data.append(["2. ความพร้อมก่อนเริ่มงาน", ""])
    for item in ready_items:
      data.append(["-", item])

    data.append(["---", "---"])
    data.append(["3. ประเด็น KYT ที่คุยกับทีมวันนี้", ""])
    for item in kyt_items:
      data.append(["-", item])

    data.append(["---", "---"])
    data.append(["4. ปัญหาค้าง/เรื่องที่ต้องประสาน", ""])
    for item in problem_items:
      data.append(["-", item])

    data.append(["---", "---"])
    data.append(["5. สถานะสาขา", status])
    data.append(["รายละเอียดสถานะ:", status_detail])

    for row_idx, row_data in enumerate(data, start=1):
      ws.cell(row=row_idx, column=1, value=row_data[0])
      ws.cell(row=row_idx, column=2, value=row_data[1])

    file_name = "Morning_KYT_Report.xlsx"
    wb.save(file_name)
    st.session_state.file_ready = True
    st.success(f"บันทึกข้อมูลสำเร็จ! ไฟล์ถูกเซฟชื่อว่า: {file_name}")

# ปุ่มดาวน์โหลดไฟล์ (อยู่นอก Form)
if st.session_state.file_ready:
  file_name = "Morning_KYT_Report.xlsx"
  try:
    with open(file_name, "rb") as file:
      st.download_button(
          label="📥 คลิกที่นี่เพื่อดาวน์โหลดไฟล์ Excel",
          data=file,
          file_name=file_name,
          mime=(
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
          ),
      )
  except FileNotFoundError:
    st.warning("กรุณากดปุ่มบันทึกข้อมูลด้านบนก่อนดาวน์โหลด")